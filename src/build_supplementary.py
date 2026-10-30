"""Builds the supplementary workbook (Tables S1 to S12) from the result files.
Run after consolidate_results.py and analyse_unseen_houses.py. Output: results_final/Supplementary_Material.xlsx"""
import os, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..")); P = lambda *a: os.path.join(ROOT, *a)
K = ["appliance", "arm", "model", "imbalance"]
first = pd.read_csv(P("results_final", "pooled_first_pass_280.csv")); final = pd.read_csv(P("results_final", "pooled_final_280.csv"))
RE = ["SVM_RBF", "RandomForest", "XGBoost", "LightGBM"]
m = first[first.model.isin(RE)].merge(final[final.model.isin(RE)], on=K, suffixes=("_first_pass", "_final"))
m["f1_change"] = m.f1_mean_final - m.f1_mean_first_pass; m["accuracy_change"] = m.accuracy_mean_final - m.accuracy_mean_first_pass
dc = pd.read_csv(P("results_depth_control", "summary_results.csv")).merge(first[first.model == "RandomForest"][K + ["f1_mean"]], on=K, suffixes=("_depth12", "_first_pass"))
dc["f1_change"] = dc.f1_mean_depth12 - dc.f1_mean_first_pass
A = json.load(open(P("results_final", "anova_pooled_270.json")))
R = json.load(open(P("results_final", "analysis_unseen_houses.json")))
anova = pd.DataFrame([("270 configurations, pooled CV, main effects", k, v) for k, v in A["main"].items()] +
                     [("270 configurations, pooled CV, two-way interactions", k, v) for k, v in A["twoway"].items()] +
                     [("270 configurations, unseen houses, main effects", k, v) for k, v in R["eta"]["lho"].items()], columns=["analysis", "term", "share_of_variance_pct"])
U = lambda ph: pd.read_csv(P("results_all_experiments", ph, "fold_results.csv"))
sheets = [
 ("Table S1", "Per-house data audit of the five appliances (share of time ON and mean power).", pd.read_csv(P("tables", "data_audit_summary.csv"))),
 ("Table S2", "Pooled nested-CV results of the first pass (280 configurations).", first),
 ("Table S3", "Pooled nested-CV final results (280 configurations).", final),
 ("Table S4", "First pass versus final run for the 160 re-run configurations.", m),
 ("Table S5", "Random Forest depth control: first-pass subsamples with the depth grid restricted to 12.", dc),
 ("Table S6", "Leave-house-out validation of the best FS-B configuration of each model (35 configurations, fold level).", pd.read_csv(P("results_leave_house_out_all", "fold_results_unified.csv"))),
 ("Table S7", "Leave-house-out factorial of LR, DT and XGBoost (120 configurations, fold level).", pd.read_csv(P("results_lho_factorial", "fold_results.csv"))),
 ("Table S8", "Share of variance in F1 by design factor (type II sums of squares).", anova),
 ("Table S9", "Leave-house-out factorial of LightGBM, RF, k-NN and SVM-RBF (fold level).", pd.concat([U("phase_A"), U("phase_E")], ignore_index=True)),
 ("Table S10", "Repeated house partitions (seeds 1 to 3), LR, DT, XGBoost and LightGBM (fold level).", U("phase_B")),
 ("Table S11", "House-grouped hyperparameter tuning, FS-B (fold level).", U("phase_C")),
 ("Table S12", "Same-protocol convolutional network on raw aggregate windows (fold level).", U("phase_D")),
]
wb = Workbook(); ws = wb.active; ws.title = "Notes"
ws["A1"] = "Supplementary material: What matters in classical non-intrusive load monitoring? A factorial benchmark on REFIT with validation on unseen houses"; ws["A1"].font = Font(bold=True)
for i, (n, t, _) in enumerate(sheets, 3): ws.cell(row=i, column=1, value=f"{n}. {t}")
ws.cell(row=len(sheets) + 4, column=1, value="F1 is the F1 score of the ON class. FS-A: basic features; FS-B: engineered rolling-window features. Class weighting is not applicable to k-NN.")
ws.column_dimensions["A"].width = 140
for n, t, df in sheets:
    s = wb.create_sheet(n); s.cell(row=1, column=1, value=f"{n}. {t}").font = Font(bold=True)
    for j, c in enumerate(df.columns, 1): s.cell(row=3, column=j, value=c).font = Font(bold=True)
    for i, row in enumerate(df.itertuples(index=False), 4):
        for j, v in enumerate(row, 1): s.cell(row=i, column=j, value=(round(v, 6) if isinstance(v, float) else v))
    s.freeze_panes = "A4"
wb.save(P("results_final", "Supplementary_Material.xlsx")); print("saved", [(n, len(d)) for n, _, d in sheets])
